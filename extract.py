import io
import re
from datetime import datetime


# ----------------------------------------------------------------- AI vision (optional upgrade)

_VISION_MEDIA_TYPES = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "gif": "image/gif", "webp": "image/webp",
}


def vision_extract_text(file_bytes, ext, api_key, model=None):
    """Ask Claude to read an image via its vision API — a stronger alternative to the
    free local OCR pipeline for busy or low-quality documents (several line items on
    one page, handwriting, glare, small print). Returns "" on anything it can't handle
    (unsupported format, no api_key, an API error) so the caller always has a clean
    signal to fall back to local OCR — this is an optional upgrade, never a dependency."""
    media_type = _VISION_MEDIA_TYPES.get((ext or "").lower())
    if not media_type or not api_key:
        return ""
    try:
        import base64
        from anthropic import Anthropic

        client = Anthropic(api_key=api_key)
        b64 = base64.standard_b64encode(file_bytes).decode("utf-8")
        resp = client.messages.create(
            model=model or "claude-sonnet-5",
            max_tokens=1024,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                    {"type": "text", "text": (
                        "Transcribe every readable piece of text in this image verbatim — all "
                        "labels, numbers, units, dates and totals — as plain text, one item per "
                        "line where possible. Do not summarize, interpret, or omit anything. If "
                        "nothing is readable, reply with nothing."
                    )},
                ],
            }],
        )
        return "".join(getattr(b, "text", "") for b in resp.content)
    except Exception:
        return ""


# ----------------------------------------------------------------- text pull

def _pdf_text(file_bytes):
    """Embedded text first; if the PDF looks scanned (little/no text),
    rasterize the first couple of pages and OCR them."""
    text = ""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            text += page.get_text() + "\n"
        if len(text.strip()) < 50:
            try:
                import pytesseract
                from PIL import Image
                ocr_text = ""
                for page in doc[:3]:
                    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3))
                    img = Image.frombytes(
                        "RGB" if pix.n < 4 else "RGBA", [pix.width, pix.height], pix.samples
                    )
                    if img.mode != "RGB":
                        img = img.convert("RGB")
                    ocr_text += pytesseract.image_to_string(img) + "\n"
                if len(ocr_text.strip()) > len(text.strip()):
                    text = ocr_text
            except Exception:
                pass
        doc.close()
    except Exception:
        pass
    return text


def _image_text(file_bytes, ext):
    try:
        from PIL import Image
        import pytesseract
        if ext == "heic":
            try:
                import pillow_heif
                pillow_heif.register_heif_opener()
            except Exception:
                pass
        img = Image.open(io.BytesIO(file_bytes))
        img = img.convert("RGB")
        return pytesseract.image_to_string(img)
    except Exception:
        return ""


def _docx_text(file_bytes):
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _xlsx_text(file_bytes):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    out.append(" ".join(cells))
        return "\n".join(out)
    except Exception:
        return ""


def _plain_text(file_bytes):
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def _eml_text(file_bytes):
    try:
        import email
        msg = email.message_from_bytes(file_bytes)
        parts = [f"Subject: {msg.get('Subject', '')}", f"Date: {msg.get('Date', '')}"]
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    payload = part.get_payload(decode=True)
                    if payload:
                        parts.append(payload.decode("utf-8", errors="ignore"))
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                parts.append(payload.decode("utf-8", errors="ignore"))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_text(file_bytes, ext):
    ext = (ext or "").lower()
    if ext == "pdf":
        return _pdf_text(file_bytes)
    if ext in ("png", "jpg", "jpeg", "gif", "webp", "heic"):
        return _image_text(file_bytes, ext)
    if ext in ("docx",):
        return _docx_text(file_bytes)
    if ext in ("xlsx", "xls"):
        return _xlsx_text(file_bytes)
    if ext in ("csv", "txt"):
        return _plain_text(file_bytes)
    if ext == "eml":
        return _eml_text(file_bytes)
    return ""  # doc / msg / unsupported binary formats: no reliable pure-python reader


# ----------------------------------------------------------------- date find

MONTHS = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "oct": 10,
    "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}

_DATE_PATTERNS = [
    (re.compile(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b"),
     lambda m: (int(m.group(1)), int(m.group(2)), int(m.group(3)))),
    (re.compile(r"\b(\d{1,2})[/.](\d{1,2})[/.](\d{4})\b"),
     lambda m: (int(m.group(3)), int(m.group(2)), int(m.group(1)))),
    (re.compile(r"\b(\d{1,2})\s+([A-Za-z]{3,9})\s+(\d{4})\b"),
     lambda m: (int(m.group(3)), MONTHS.get(m.group(2).lower(), 0), int(m.group(1)))),
    (re.compile(r"\b([A-Za-z]{3,9})\s+(\d{1,2}),?\s+(\d{4})\b"),
     lambda m: (int(m.group(3)), MONTHS.get(m.group(1).lower(), 0), int(m.group(2)))),
]

_DATE_KEYWORDS = ["date", "period", "billing", "invoice", "statement", "reading", "issued", "due"]


def find_date(text):
    found = []
    for pattern, parser in _DATE_PATTERNS:
        for m in pattern.finditer(text):
            try:
                y, mo, d = parser(m)
                if not (1 <= mo <= 12) or not (1 <= d <= 31) or not (2000 <= y <= 2100):
                    continue
                dt = datetime(y, mo, d)
            except Exception:
                continue
            line_start = text.rfind("\n", 0, m.start()) + 1
            line_end = text.find("\n", m.end())
            if line_end == -1:
                line_end = len(text)
            context = text[line_start:line_end].lower()
            score = 1 + sum(1 for kw in _DATE_KEYWORDS if kw in context)
            found.append((score, m.start(), dt))
    if not found:
        return None
    found.sort(key=lambda x: (-x[0], x[1]))
    return found[0][2].strftime("%Y-%m-%d")


# ---------------------------------------------------------------- value find

# Grouped-thousands alternative requires at least one real separator group (+, not *) so an
# ungrouped number like "2026" falls through to the plain-digit alternative instead of being
# truncated to its first 1-3 digits (regex alternation picks the first alternative that matches
# at all, not the longest — a `*` here previously let "2026" wrongly match as just "202").
_NUMBER = r"([-+]?\d{1,3}(?:[ ,]\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"
_VALUE_KEYWORDS = ["total", "consumption", "usage", "amount due", "amount", "net", "reading", "quantity", "volume"]


# Unit sub-tokens that label the KIND of value expected (a date, a status word, a name)
# rather than a quantity — matching a stray number near the word "date" or "status" itself
# (e.g. a year inside a nearby date) produces a nonsense "value", so these are skipped as
# match targets. The real date is already captured separately by find_date().
_PLACEHOLDER_UNIT_TOKENS = {"date", "status", "name", "topic", "y/n", "yes", "no", "level", "y", "n"}


def find_value(text, unit_field):
    units = [u.strip() for u in re.split(r"[,/]", unit_field or "") if u.strip()]

    # 1. number directly touching a known unit token -> high confidence
    for unit in units:
        if unit.lower() in _PLACEHOLDER_UNIT_TOKENS:
            continue
        u = re.escape(unit)
        m = re.search(_NUMBER + r"\s*" + u + r"\b", text, re.IGNORECASE)
        if m:
            return f"{m.group(1).strip()} {unit}", "high"
        m = re.search(u + r"\s*[:\-]?\s*" + _NUMBER, text, re.IGNORECASE)
        if m:
            return f"{m.group(1).strip()} {unit}", "high"

    # 2. percentages
    if any(u == "%" for u in units):
        m = re.search(_NUMBER + r"\s*%", text)
        if m:
            return f"{m.group(1).strip()}%", "high"

    # 3. a number next to a generic total/usage-style keyword -> medium confidence
    for kw in _VALUE_KEYWORDS:
        m = re.search(re.escape(kw) + r"[:\s]{1,15}" + _NUMBER, text, re.IGNORECASE)
        if m:
            hint = next((u for u in units if u.lower() not in ("yes", "no")), "")
            return f"{m.group(1).strip()}{(' ' + hint) if hint else ''}", "medium"

    return None, None


# --------------------------------------------------------------------- notes

def build_notes(text, source_name):
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return None
    snippet = cleaned[:280] + ("…" if len(cleaned) > 280 else "")
    return f'Auto-extracted from {source_name}: "{snippet}"'


# ------------------------------------------------------------ single-item analysis

def analyze_text(text, source_name, item):
    """Core analysis against ONE checklist item (dict-like row, needs 'unit'), or
    None for a date/notes-only read with no value matching."""
    if not text or not text.strip():
        return {
            "ok": False,
            "message": "Couldn't read anything from this — please fill in the fields manually.",
            "value_text": None,
            "value_confidence": None,
            "entry_date": None,
            "notes": None,
            "raw_text": "",
        }

    entry_date = find_date(text)
    unit_field = (item or {}).get("unit") if item else None
    value_text, confidence = find_value(text, unit_field)
    notes = build_notes(text, source_name)

    return {
        "ok": True,
        "value_text": value_text,
        "value_confidence": confidence,
        "entry_date": entry_date,
        "notes": notes,
        "raw_text": text[:4000],
    }


def extract_evidence(file_bytes, filename, item):
    """Back-compat wrapper: read a file, then analyze it against one item."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        text = extract_text(file_bytes, ext)
    except Exception:
        text = ""
    return analyze_text(text, filename, item)


# ------------------------------------------------------------- multi-item analysis

_KEYWORD_STOPWORDS = {
    "and", "the", "of", "by", "or", "per", "in", "on", "for", "with", "to", "a", "an",
    "from", "that", "this", "their", "number", "total", "share", "level", "was", "are",
    "its", "has", "had", "not", "but", "you", "all", "any", "via", "into", "each",
}


def _item_keywords(item):
    # {3,} (not {4,}) so short but meaningful acronyms like "LPG" still count as keywords —
    # the expanded stopword list above keeps this from pulling in ordinary short words.
    words = re.findall(r"[a-zA-Z]{3,}", f"{item.get('data_point','')} {item.get('category','')}")
    return {w.lower() for w in words if w.lower() not in _KEYWORD_STOPWORDS}


def analyze_text_multi(text, source_name, items):
    """One photo/note -> many fields: scan the same text against every checklist
    item and return every plausible match as a candidate entry, instead of
    requiring the person to pick a single data point up front. Each candidate is
    a suggestion for the person to confirm, edit or discard, never auto-saved.

    To avoid every same-unit item (e.g. a dozen different "kg" data points)
    flooding the review list with the same one number, a match is only
    "high confidence" when it comes from a line/sentence that also contains one
    of that item's own keywords (drawn from its data_point + category text) —
    so "LPG used: 18 kg" anchors to the LPG item, not to every kg-denominated
    item in the checklist. When an item's unit is unique across the whole
    checklist there's no ambiguity to anchor against, so a whole-text match is
    still allowed, just demoted to medium confidence. Deliberately capped at 12
    candidates so a very generic document doesn't overwhelm the confirmation
    screen."""
    if not text or not text.strip():
        return {
            "ok": False,
            "message": "Couldn't read anything from this — please try a clearer photo, or enter it manually.",
            "entry_date": None,
            "notes": None,
            "candidates": [],
            "raw_text": "",
        }

    entry_date = find_date(text)
    notes = build_notes(text, source_name)

    # Split on newlines and sentence-ending periods (a period followed by whitespace),
    # never a bare "." — splitting on every "." would cut a decimal value like "42.5"
    # into "42" and "5" as separate lines.
    lines = [l for l in re.split(r"\n|\.\s+", text) if l.strip()]
    unit_counts = {}
    for item in items:
        unit_counts[item.get("unit")] = unit_counts.get(item.get("unit"), 0) + 1

    candidates = []
    for item in items:
        kws = _item_keywords(item)
        value_text, confidence = None, None
        for line in lines:
            if kws and any(kw in line.lower() for kw in kws):
                value_text, confidence = find_value(line, item.get("unit"))
                if value_text:
                    break
        if not value_text and unit_counts.get(item.get("unit"), 0) == 1:
            value_text, confidence = find_value(text, item.get("unit"))
            if confidence == "high":
                confidence = "medium"  # unanchored — demote so it doesn't outrank a real keyword match
        if not value_text:
            continue
        candidates.append({
            "item_id": item["id"],
            "data_point": item["data_point"],
            "pillar": item["pillar"],
            "category": item["category"],
            "unit": item.get("unit"),
            "value_text": value_text,
            "value_confidence": confidence,
        })

    # highest-confidence matches first, capped so the review list stays short
    order = {"high": 0, "medium": 1, None: 2}
    candidates.sort(key=lambda c: order.get(c["value_confidence"], 2))
    candidates = candidates[:12]

    return {
        "ok": bool(candidates),
        "message": None if candidates else
            "Read the text but couldn't confidently match it to a data point — please enter it manually.",
        "entry_date": entry_date,
        "notes": notes,
        "candidates": candidates,
        "raw_text": text[:4000],
    }
